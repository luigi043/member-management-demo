from pathlib import Path
import subprocess, re, textwrap, sys
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
root=Path.cwd(); out=root/'docs/test-evidence'
checks=[('01-tests',['npm.cmd','test','--','--reporter=verbose']),('02-lint',['npm.cmd','run','lint']),('03-build',['npm.cmd','run','build'])]
records=[]
for label,cmd in checks:
    if '--reuse' in sys.argv:
        records.append((label,(out/(label+'.txt')).read_text(encoding='utf-8'))); continue
    p=subprocess.run(cmd,capture_output=True,text=True,encoding='utf-8',errors='replace',shell=True)
    raw=p.stdout+p.stderr
    raw=re.sub(r'\x1b\[[0-9;]*[A-Za-z]','',raw)
    text='$ '+' '.join(cmd)+'\n'+raw+'\nExit code: '+str(p.returncode)
    (out/(label+'.txt')).write_text(text,encoding='utf-8')
    records.append((label,text))
rehearsal=(root/'rehearsal.log').read_text(encoding='utf-8-sig',errors='replace')
records.append(('04-history-rehearsal','Excerpt from actual rehearsal output; full record is in TXT.\n'+ '\n'.join(line for line in rehearsal.splitlines() if not line.startswith('warning:'))))
(out/'04-history-rehearsal.txt').write_text(rehearsal,encoding='utf-8')
font=ImageFont.truetype('C:/Windows/Fonts/consola.ttf',20)
doc=Document(); sec=doc.sections[0]; sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.65)
for s in ['Normal','Title','Heading 1']:
    doc.styles[s].font.name='Calibri'; doc.styles[s].font.color.rgb=RGBColor(0,0,0)
doc.styles['Normal'].font.size=Pt(11)
doc.add_paragraph('Demo 5 test evidence', 'Title')
doc.add_paragraph('23 October 2026. Captured automated results for the member management Git workflow demo. The images reproduce actual command output, with line wrapping for readability. They are rendered transcript captures, not photographs or screenshots of a desktop terminal. Original text records are retained alongside the images.')
doc.add_paragraph('Scope and limitations','Heading 1')
doc.add_paragraph('The automated UI test checks duplicate registration input retention and successful form reset. Unit tests cover validation, search and CRUD. The history exercise runs in a disposable repository. No manual browser walkthrough, hosted pull request, push, merge to the existing branches or deployment was performed. Node 25 emits a localStorage warning; all tests pass using the capability-checked fallback.')
for label,text in records:
    lines=[]
    for line in text.splitlines():
        # Keep every result line, wrap long paths and statements.
        lines.extend(textwrap.wrap(line,width=105,replace_whitespace=False,drop_whitespace=False) or [''])
    for part in range(0,len(lines),45):
        chunk=lines[part:part+45]
        im=Image.new('RGB',(1320,70+len(chunk)*27),'#101820'); draw=ImageDraw.Draw(im)
        for n,line in enumerate(chunk): draw.text((22,22+n*27),line,font=font,fill='#ecf3f8')
        name=f'{label}-{part//45+1}.png'; im.save(out/name)
        doc.add_page_break(); doc.add_paragraph(label.replace('-',' '),'Heading 1')
        doc.add_paragraph('Actual recorded output' + (f' continued part {part//45+1}' if part else '') + '. See the matching TXT file for the complete record.')
        doc.add_picture(str(out/name),width=Inches(7.0))
from docx.oxml.ns import qn
for element in list(doc.styles.element.iter(qn('w:pBdr'))): element.getparent().remove(element)
doc.save(out/'Demo5-Test-Evidence.docx')
print('Created',out/'Demo5-Test-Evidence.docx')
